"""Transaction template / import / export with unit isolation."""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from adapters.api.deps import get_current_user, require_roles
from adapters.api.scope import (
    WRITE_ROLES,
    assert_can_mutate_period,
    assert_not_readonly,
    is_pengelola,
    is_readonly,
    scoped_unit_id,
    unit_code_for,
)
from adapters.external.excel_adapter import parse_excel_rows
from modules.identity.infrastructure.models import User
from modules.siabumdes.application.services import FinanceService
from modules.siabumdes.infrastructure.models import Account, Transaction, TransactionType, UnitUsaha
from shared.config import REPORT_READ_LEVEL, public_role
from shared.database import get_db

router = APIRouter(prefix="/api", tags=["import-export"])

TEMPLATE_HEADERS = ["tanggal", "unit_code", "jenis_transaksi", "keterangan", "nominal", "debit", "kredit", "referensi"]


def _xlsx(buf: BytesIO, filename: str) -> Response:
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/transactions/template")
@router.get("/import/template")
async def transaction_template(
    user: User = Depends(require_roles(*WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Transaksi"
    ws.append(TEMPLATE_HEADERS)
    sample_unit = "UU01"
    if is_pengelola(user):
        unit = await session.get(UnitUsaha, user.unit_usaha_id)
        sample_unit = unit.code if unit else "UU01"
    ws.append(["2026-01-15", sample_unit, "contoh_transaksi", "Contoh baris import", 100000, "", "", "nota-001"])
    ref = wb.create_sheet("Referensi")
    ref.append(["Kode Unit", "Nama Unit"])
    units = (await session.execute(select(UnitUsaha).order_by(UnitUsaha.code))).scalars()
    for u in units:
        if is_pengelola(user) and u.id != user.unit_usaha_id:
            continue
        ref.append([u.code, u.name])
    ref.append([])
    ref.append(["Kode Jenis Transaksi", "Nama", "Group"])
    types = (await session.execute(select(TransactionType))).scalars()
    for t in types:
        if is_pengelola(user):
            unit = await session.get(UnitUsaha, user.unit_usaha_id)
            if unit and t.group_code not in {unit.code, "BUMDES"}:
                continue
        ref.append([t.code, t.name, t.group_code])
    buf = BytesIO()
    wb.save(buf)
    return _xlsx(buf, "Template-Transaksi-BUMDES.xlsx")


@router.post("/transactions/import")
@router.post("/import")
async def import_transactions(
    file: UploadFile = File(...),
    user: User = Depends(require_roles(*WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    assert_not_readonly(user)
    raw = await file.read()
    try:
        wb = load_workbook(BytesIO(raw), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"File Excel tidak valid: {exc}") from exc
    ws = wb.active
    header = [str(c.value or "").strip().lower() for c in ws[1]]

    def col(name: str) -> Optional[int]:
        try:
            return header.index(name)
        except ValueError:
            return None

    idx_tanggal, idx_type, idx_amt = col("tanggal"), col("jenis_transaksi"), col("nominal")
    if None in {idx_tanggal, idx_type, idx_amt}:
        raise HTTPException(status_code=400, detail="Kolom wajib: tanggal, jenis_transaksi, nominal")
    idx_unit, idx_desc, idx_deb, idx_kre, idx_ref = col("unit_code"), col("keterangan"), col("debit"), col("kredit"), col("referensi")

    types = {(t.code, t.group_code): t for t in (await session.execute(select(TransactionType))).scalars()}
    types_by_code = {}
    for t in types.values():
        types_by_code.setdefault(t.code, t)
    units = {u.code: u for u in (await session.execute(select(UnitUsaha))).scalars()}
    accounts = {(a.code, a.group_code): a for a in (await session.execute(select(Account))).scalars()}

    inserted = 0
    errors: list[dict] = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v in (None, "") for v in row):
            continue
        try:
            tanggal = row[idx_tanggal]
            if hasattr(tanggal, "strftime"):
                tanggal = tanggal.strftime("%Y-%m-%d")
            tx_date = date.fromisoformat(str(tanggal).strip()[:10])
            type_code = str(row[idx_type] or "").strip()
            tt = types_by_code.get(type_code)
            if not tt:
                raise ValueError(f"Jenis transaksi '{type_code}' tidak dikenal")
            amount = Decimal(str(row[idx_amt] or 0)).quantize(Decimal("0.01"))
            if amount <= 0:
                raise ValueError("Nominal harus > 0")
            unit_code = str(row[idx_unit] or "").strip().upper() if idx_unit is not None else ""
            unit = units.get(unit_code) if unit_code else None
            unit_id = unit.id if unit else None
            if is_pengelola(user):
                unit_id = user.unit_usaha_id
                own = await session.get(UnitUsaha, user.unit_usaha_id)
                if unit_code and own and unit_code != own.code:
                    raise ValueError("Pengelola hanya boleh import data unit sendiri")
            await assert_can_mutate_period(session, user, tx_date, unit_id)
            group = await unit_code_for(session, unit_id)
            debit = str(row[idx_deb] or "").strip() if idx_deb is not None else tt.debit
            credit = str(row[idx_kre] or "").strip() if idx_kre is not None else tt.credit
            debit = debit or tt.debit
            credit = credit or tt.credit
            if (debit, group) not in accounts or (credit, group) not in accounts:
                raise ValueError(f"Akun {debit}/{credit} tidak ada di COA {group}")
            desc = str(row[idx_desc]).strip() if idx_desc is not None and row[idx_desc] else tt.name
            ref = str(row[idx_ref]).strip() if idx_ref is not None and row[idx_ref] else ""
            tx = Transaction(
                date=tx_date,
                unit_usaha_id=unit_id,
                transaction_type=type_code,
                description=desc,
                amount=amount,
                debit_account_code=debit,
                credit_account_code=credit,
                reference=ref,
                created_by=user.id,
                proofs=[],
            )
            session.add(tx)
            await session.flush()
            await FinanceService(session).create_journal_entry(
                transaction_id=tx.id,
                entry_date=tx_date,
                memo=desc,
                debit_account_id=accounts[(debit, group)].id,
                credit_account_id=accounts[(credit, group)].id,
                amount=amount,
            )
            inserted += 1
        except Exception as exc:
            errors.append({"row": row_num, "error": str(exc)})
    return {"inserted": inserted, "errors": errors, "total_rows": max(ws.max_row - 1, 0)}


@router.get("/transactions/export")
@router.get("/export")
async def export_transactions(
    unit_usaha_id: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    all_data: bool = Query(False),
    user: User = Depends(require_roles(*REPORT_READ_LEVEL)),
    session: AsyncSession = Depends(get_db),
):
    units = {u.id: u for u in (await session.execute(select(UnitUsaha))).scalars()}
    accounts = {a.code: a for a in (await session.execute(select(Account))).scalars()}
    stmt = select(Transaction).order_by(Transaction.date.asc())
    if is_pengelola(user):
        stmt = stmt.where(Transaction.unit_usaha_id == user.unit_usaha_id)
    elif not all_data:
        if unit_usaha_id == "":
            stmt = stmt.where(Transaction.unit_usaha_id.is_(None))
        elif unit_usaha_id:
            stmt = stmt.where(Transaction.unit_usaha_id == unit_usaha_id)
        if start_date:
            stmt = stmt.where(Transaction.date >= date.fromisoformat(start_date[:10]))
        if end_date:
            stmt = stmt.where(Transaction.date <= date.fromisoformat(end_date[:10]))
    txs = list((await session.execute(stmt)).scalars())

    headers = [
        "Tanggal", "Kelompok", "Jenis Transaksi", "Keterangan",
        "Kode Debit", "Nama Debit", "Kode Kredit", "Nama Kredit", "Nominal", "Referensi",
    ]

    def row_of(tx: Transaction) -> list:
        grp = units[tx.unit_usaha_id].code if tx.unit_usaha_id and tx.unit_usaha_id in units else "BUMDES"
        return [
            tx.date.isoformat(),
            grp,
            tx.transaction_type,
            tx.description,
            tx.debit_account_code,
            (accounts.get(tx.debit_account_code) or Account(code="", name="")).name if False else (accounts.get(tx.debit_account_code).name if accounts.get(tx.debit_account_code) else ""),
            tx.credit_account_code,
            accounts.get(tx.credit_account_code).name if accounts.get(tx.credit_account_code) else "",
            float(tx.amount),
            tx.reference or "",
        ]

    wb = Workbook()
    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")

    def fill(ws, title: str, items: list[Transaction]):
        ws.append([title])
        ws.append(headers)
        for cell in ws[2]:
            cell.font = header_font
            cell.fill = header_fill
        for tx in items:
            ws.append(row_of(tx))

    if all_data and not is_pengelola(user):
        ws0 = wb.active
        ws0.title = "Semua"
        fill(ws0, "Semua Transaksi", txs)
        ws_b = wb.create_sheet("BUMDES")
        fill(ws_b, "BUMDES", [t for t in txs if not t.unit_usaha_id])
        for u in sorted(units.values(), key=lambda x: x.code):
            ws_u = wb.create_sheet(u.code)
            fill(ws_u, u.code, [t for t in txs if t.unit_usaha_id == u.id])
        fname = f"Transaksi_Semua_Data_{datetime.now().strftime('%Y%m%d')}.xlsx"
    else:
        ws = wb.active
        ws.title = "Transaksi"
        label = "Unit" if is_pengelola(user) else ("BUMDES" if unit_usaha_id == "" else "Filter")
        fill(ws, f"Transaksi {label}", txs)
        fname = f"Transaksi_{label}.xlsx"
    buf = BytesIO()
    wb.save(buf)
    return _xlsx(buf, fname)
