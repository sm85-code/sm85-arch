"""Master data HTTP API matching frontend-siabumdes paths."""
from __future__ import annotations

from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from adapters.api.deps import get_current_user, require_roles
from adapters.api.scope import (
    MASTER_WRITE_ROLES,
    UNIT_WRITE_ROLES,
    is_pengelola,
    pengelola_unit,
    scoped_unit_id,
)
from adapters.external.excel_adapter import generate_excel_report, parse_excel_rows
from modules.identity.infrastructure.models import User
from modules.siabumdes.infrastructure.models import Account, Mitra, TransactionType, UnitUsaha
from shared.database import get_db

router = APIRouter(prefix="/api", tags=["master-data"])

VALID_CATEGORIES = {
    "aset": {"kas_bank", "piutang", "persediaan", "aset_tetap", "aset_lain"},
    "kewajiban": {"utang_usaha", "utang_lain", "utang_pajak"},
    "ekuitas": {"modal", "saldo_laba"},
    "pendapatan": {"pendapatan_usaha", "pendapatan_lain"},
    "hpp": {"hpp"},
    "beban": {"beban_operasi", "beban_lain"},
}


class AccountIn(BaseModel):
    code: str
    name: str
    category: str
    subcategory: str = ""
    normal_balance: str
    group: str = "BUMDES"


class UnitIn(BaseModel):
    code: str
    name: str
    description: str = ""
    revenue_scheme: str = ""


class TxTypeIn(BaseModel):
    code: str
    name: str
    debit: str
    credit: str
    unit_codes: list[str] = Field(default_factory=list)
    group: str = "BUMDES"


class MitraIn(BaseModel):
    name: str
    unit_usaha_id: Optional[str] = None
    phone: str = ""
    note: str = ""


def _acc_out(row: Account) -> dict:
    return {
        "id": row.id,
        "code": row.code,
        "name": row.name,
        "category": row.category,
        "subcategory": row.subcategory,
        "normal_balance": row.normal_balance,
        "group": row.group_code,
        "group_code": row.group_code,
        "unit_usaha_id": row.unit_usaha_id,
        "active": row.active,
    }


def _type_out(row: TransactionType) -> dict:
    return {
        "id": row.id,
        "code": row.code,
        "name": row.name,
        "debit": row.debit,
        "credit": row.credit,
        "group": row.group_code,
        "unit_codes": list(row.unit_codes or []),
    }


async def _group_filter(session: AsyncSession, user: User) -> Optional[str]:
    if is_pengelola(user):
        return (await pengelola_unit(session, user)).code
    return None


@router.get("/unit-usaha")
async def list_units(user: User = Depends(get_current_user), session: AsyncSession = Depends(get_db)):
    stmt = select(UnitUsaha).order_by(UnitUsaha.code.asc())
    if is_pengelola(user):
        stmt = stmt.where(UnitUsaha.id == user.unit_usaha_id)
    rows = (await session.execute(stmt)).scalars()
    return [
        {
            "id": row.id,
            "code": row.code,
            "name": row.name,
            "description": row.description,
            "revenue_scheme": row.revenue_scheme,
            "active": row.active,
        }
        for row in rows
    ]


@router.post("/unit-usaha")
async def create_unit(
    payload: UnitIn,
    _: User = Depends(require_roles(*UNIT_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    exists = (await session.execute(select(UnitUsaha).where(UnitUsaha.code == payload.code.strip()))).scalar_one_or_none()
    if exists:
        raise HTTPException(status_code=400, detail="Kode unit sudah ada")
    row = UnitUsaha(
        code=payload.code.strip().upper(),
        name=payload.name.strip(),
        description=payload.description,
        revenue_scheme=payload.revenue_scheme,
    )
    session.add(row)
    await session.flush()
    return {"id": row.id, "code": row.code, "name": row.name, "description": row.description, "revenue_scheme": row.revenue_scheme}


@router.get("/accounts")
async def list_accounts(user: User = Depends(get_current_user), session: AsyncSession = Depends(get_db)):
    stmt = select(Account).order_by(Account.group_code.asc(), Account.code.asc())
    forced = await _group_filter(session, user)
    if forced:
        stmt = stmt.where(Account.group_code == forced)
    rows = (await session.execute(stmt)).scalars()
    return [_acc_out(row) for row in rows]


@router.post("/accounts")
async def create_account(
    payload: AccountIn,
    _: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    group = (payload.group or "BUMDES").strip().upper()
    exists = (
        await session.execute(select(Account).where(Account.code == payload.code, Account.group_code == group))
    ).scalar_one_or_none()
    if exists:
        raise HTTPException(status_code=400, detail=f"Kode akun {payload.code} sudah ada di kelompok {group}")
    if payload.category not in VALID_CATEGORIES:
        raise HTTPException(status_code=400, detail="Kategori tidak valid")
    if payload.normal_balance not in {"debit", "kredit"}:
        raise HTTPException(status_code=400, detail="normal_balance harus debit atau kredit")
    unit = (await session.execute(select(UnitUsaha).where(UnitUsaha.code == group))).scalar_one_or_none()
    row = Account(
        code=payload.code.strip(),
        name=payload.name.strip(),
        category=payload.category.strip().lower(),
        subcategory=(payload.subcategory or "").strip().lower(),
        normal_balance=payload.normal_balance,
        group_code=group,
        unit_usaha_id=unit.id if unit else None,
    )
    session.add(row)
    await session.flush()
    return _acc_out(row)


@router.put("/accounts/{code}")
async def update_account(
    code: str,
    payload: AccountIn,
    group: str = Query("BUMDES"),
    _: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    row = (
        await session.execute(select(Account).where(Account.code == code, Account.group_code == group))
    ).scalar_one_or_none()
    if not row:
        raise HTTPException(status_code=404, detail=f"Kode akun {code} tidak ditemukan di kelompok {group}")
    target_group = (payload.group or group).strip().upper()
    row.code = payload.code.strip()
    row.name = payload.name.strip()
    row.category = payload.category.strip().lower()
    row.subcategory = (payload.subcategory or "").strip().lower()
    row.normal_balance = payload.normal_balance
    row.group_code = target_group
    await session.flush()
    return _acc_out(row)


@router.delete("/accounts/{code}")
async def delete_account(
    code: str,
    group: str = Query("BUMDES"),
    _: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    row = (
        await session.execute(select(Account).where(Account.code == code, Account.group_code == group))
    ).scalar_one_or_none()
    if not row:
        raise HTTPException(status_code=404, detail="Kode akun tidak ditemukan")
    await session.delete(row)
    return {"deleted": 1}


@router.delete("/accounts/reset-all")
async def reset_accounts(
    confirm: str = "",
    _: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    if confirm != "YES":
        raise HTTPException(status_code=400, detail="Parameter confirm=YES wajib untuk operasi ini")
    rows = (await session.execute(select(Account))).scalars().all()
    for row in rows:
        await session.delete(row)
    return {"deleted": len(rows)}


@router.get("/transaction-types")
async def list_tx_types(user: User = Depends(get_current_user), session: AsyncSession = Depends(get_db)):
    stmt = select(TransactionType).order_by(TransactionType.group_code.asc(), TransactionType.code.asc())
    forced = await _group_filter(session, user)
    if forced:
        stmt = stmt.where(TransactionType.group_code == forced)
    rows = (await session.execute(stmt)).scalars()
    return [_type_out(row) for row in rows]


@router.post("/transaction-types")
async def create_tx_type(
    payload: TxTypeIn,
    _: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    group = (payload.group or "BUMDES").strip().upper()
    exists = (
        await session.execute(
            select(TransactionType).where(TransactionType.code == payload.code, TransactionType.group_code == group)
        )
    ).scalar_one_or_none()
    if exists:
        raise HTTPException(status_code=400, detail="Kode jenis transaksi sudah ada")
    row = TransactionType(
        code=payload.code.strip(),
        name=payload.name.strip(),
        debit=payload.debit,
        credit=payload.credit,
        group_code=group,
        unit_codes=payload.unit_codes or [],
    )
    session.add(row)
    await session.flush()
    return _type_out(row)


@router.put("/transaction-types/{code}")
async def update_tx_type(
    code: str,
    payload: TxTypeIn,
    _: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    group = (payload.group or "BUMDES").strip().upper()
    row = (
        await session.execute(
            select(TransactionType).where(TransactionType.code == code, TransactionType.group_code == group)
        )
    ).scalar_one_or_none()
    if not row:
        row = (await session.execute(select(TransactionType).where(TransactionType.code == code))).scalar_one_or_none()
    if not row:
        raise HTTPException(status_code=404, detail="Jenis transaksi tidak ditemukan")
    if payload.code != code:
        clash = (
            await session.execute(
                select(TransactionType).where(
                    TransactionType.code == payload.code, TransactionType.group_code == group
                )
            )
        ).scalar_one_or_none()
        if clash:
            raise HTTPException(status_code=400, detail="Kode tujuan sudah ada")
        row.code = payload.code
    row.name = payload.name
    row.debit = payload.debit
    row.credit = payload.credit
    row.group_code = group
    row.unit_codes = payload.unit_codes or []
    await session.flush()
    return _type_out(row)


@router.delete("/transaction-types/{code}")
async def delete_tx_type(
    code: str,
    _: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    row = (await session.execute(select(TransactionType).where(TransactionType.code == code))).scalar_one_or_none()
    if not row:
        raise HTTPException(status_code=404, detail="Jenis transaksi tidak ditemukan")
    await session.delete(row)
    return {"deleted": 1}


@router.get("/mitra")
async def list_mitra(
    unit_usaha_id: Optional[str] = None,
    user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_db),
):
    stmt = select(Mitra)
    if is_pengelola(user):
        stmt = stmt.where(Mitra.unit_usaha_id == user.unit_usaha_id)
    elif unit_usaha_id:
        stmt = stmt.where(Mitra.unit_usaha_id == unit_usaha_id)
    rows = (await session.execute(stmt)).scalars()
    return [
        {"id": row.id, "name": row.name, "unit_usaha_id": row.unit_usaha_id, "phone": row.phone, "note": row.note}
        for row in rows
    ]


@router.post("/mitra")
async def create_mitra(
    payload: MitraIn,
    user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_db),
):
    if user.role in {"pengawas", "penasihat"}:
        raise HTTPException(status_code=403, detail="Akun bersifat read-only")
    unit_id = await scoped_unit_id(session, user, payload.unit_usaha_id)
    row = Mitra(name=payload.name.strip(), unit_usaha_id=unit_id, phone=payload.phone, note=payload.note)
    session.add(row)
    await session.flush()
    return {"id": row.id, "name": row.name, "unit_usaha_id": row.unit_usaha_id, "phone": row.phone, "note": row.note}


@router.delete("/mitra/{mitra_id}")
async def delete_mitra(
    mitra_id: str,
    _: User = Depends(require_roles("admin", "direktur", "bendahara")),
    session: AsyncSession = Depends(get_db),
):
    row = await session.get(Mitra, mitra_id)
    if not row:
        return {"deleted": 0}
    await session.delete(row)
    return {"deleted": 1}


@router.post("/accounts/import")
async def import_accounts(
    file: UploadFile = File(...),
    _: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    content = await file.read()
    from openpyxl import load_workbook
    import io as _io

    try:
        wb = load_workbook(filename=_io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"File tidak valid: {exc}") from exc

    inserted = skipped = 0
    errors: list[str] = []
    valid_groups = {"BUMDES", "UU01", "UU02", "UU03", "UU04", "UU05", "UU06"}
    for sheet_name in wb.sheetnames:
        if sheet_name not in valid_groups:
            continue
        headers, data = parse_excel_rows(content, sheet_name=sheet_name)
        header = [h.strip().lower() for h in headers]
        if header[:5] != ["code", "name", "category", "subcategory", "normal_balance"]:
            errors.append(f"Sheet '{sheet_name}': header tidak sesuai")
            continue
        unit = (await session.execute(select(UnitUsaha).where(UnitUsaha.code == sheet_name))).scalar_one_or_none()
        for idx, values in enumerate(data, start=2):
            code = str(values[0] or "").strip()
            name = str(values[1] or "").strip()
            category = str(values[2] or "").strip().lower()
            subcategory = str(values[3] or "").strip().lower()
            normal_balance = str(values[4] or "").strip().lower()
            if not (code and name and category and normal_balance):
                skipped += 1
                continue
            exists = (
                await session.execute(select(Account).where(Account.code == code, Account.group_code == sheet_name))
            ).scalar_one_or_none()
            if exists:
                skipped += 1
                continue
            session.add(
                Account(
                    code=code,
                    name=name,
                    category=category,
                    subcategory=subcategory,
                    normal_balance=normal_balance,
                    group_code=sheet_name,
                    unit_usaha_id=unit.id if unit else None,
                )
            )
            inserted += 1
    return {"inserted": inserted, "skipped": skipped, "errors": errors[:50]}


@router.get("/accounts/template")
async def accounts_template(_: User = Depends(require_roles(*MASTER_WRITE_ROLES))):
    headers = ["code", "name", "category", "subcategory", "normal_balance"]
    rows = [["1.1.01.01", "Kas Bendahara", "aset", "kas_bank", "debit"]]
    blob = generate_excel_report(headers, rows, "BUMDES")
    return StreamingResponse(
        iter([blob]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Template-Kode-Akun.xlsx"'},
    )


@router.post("/transaction-types/import")
async def import_tx_types(
    file: UploadFile = File(...),
    _: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    content = await file.read()
    headers, data = parse_excel_rows(content)
    header = [h.strip().lower() for h in headers]
    required = {"code", "name", "debit", "credit"}
    if not required.issubset(set(header)):
        raise HTTPException(status_code=400, detail="Kolom wajib: code, name, debit, credit")
    inserted = skipped = 0
    for values in data:
        item = dict(zip(header, values))
        code = str(item.get("code") or "").strip()
        name = str(item.get("name") or "").strip()
        if not code or not name:
            continue
        group = str(item.get("group") or "BUMDES").strip().upper()
        exists = (
            await session.execute(
                select(TransactionType).where(TransactionType.code == code, TransactionType.group_code == group)
            )
        ).scalar_one_or_none()
        if exists:
            skipped += 1
            continue
        session.add(
            TransactionType(
                code=code,
                name=name,
                debit=str(item.get("debit") or "").strip(),
                credit=str(item.get("credit") or "").strip(),
                group_code=group,
            )
        )
        inserted += 1
    return {"inserted": inserted, "skipped": skipped, "errors": []}


@router.get("/transaction-types/template")
async def tx_types_template(
    group: str = Query("BUMDES"),
    _: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
):
    headers = ["code", "name", "debit", "credit", "group"]
    rows = [["contoh_transaksi", "Contoh Transaksi", "1.1.01.01", "4.1.01.01", group]]
    blob = generate_excel_report(headers, rows, "Jenis Transaksi")
    return StreamingResponse(
        iter([blob]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Template-Jenis-Transaksi.xlsx"'},
    )


@router.get("/master-data/export")
async def export_master(
    group: str = Query("BUMDES"),
    section: str = Query("all"),
    _: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    if section in {"all", "accounts"}:
        accs = (
            await session.execute(select(Account).where(Account.group_code == group).order_by(Account.code))
        ).scalars()
        headers = ["code", "name", "category", "subcategory", "normal_balance", "group"]
        rows = [[a.code, a.name, a.category, a.subcategory, a.normal_balance, a.group_code] for a in accs]
        blob = generate_excel_report(headers, rows, "Kode Akun")
        filename = f"Master-Data-Kode-Akun-{group}.xlsx"
    else:
        types = (
            await session.execute(
                select(TransactionType).where(TransactionType.group_code == group).order_by(TransactionType.code)
            )
        ).scalars()
        headers = ["code", "name", "debit", "credit", "group"]
        rows = [[t.code, t.name, t.debit, t.credit, t.group_code] for t in types]
        blob = generate_excel_report(headers, rows, "Jenis Transaksi")
        filename = f"Master-Data-Jenis-Transaksi-{group}.xlsx"
    return StreamingResponse(
        iter([blob]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
