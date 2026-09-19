"""OpenPyXL inbound parse + audit-ready outbound report generator."""
from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal
from typing import Any, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from shared.report_branding import get_report_branding


def parse_excel_rows(
    file_bytes: bytes,
    *,
    sheet_name: Optional[str] = None,
    header_row: int = 1,
) -> tuple[list[str], list[list[Any]]]:
    """Return (headers, data rows) from .xlsx bytes."""
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers: list[str] = []
    data: list[list[Any]] = []
    for idx, row in enumerate(rows_iter, start=1):
        values = list(row)
        if idx < header_row:
            continue
        if idx == header_row:
            headers = [str(c).strip() if c is not None else "" for c in values]
            continue
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        data.append(list(values))
    wb.close()
    return headers, data


# Back-compat alias used by older imports
def parse_excel_file(file_bytes: bytes, **kwargs):
    return parse_excel_rows(file_bytes, **kwargs)


def _is_money_header(h: str) -> bool:
    h = (h or "").lower()
    return any(
        k in h
        for k in (
            "nominal", "jumlah", "debit", "kredit", "saldo", "rp",
            "amount", "nilai", "laba", "beban", "pendapatan",
        )
    )


def _row_is_total(row: list[Any]) -> bool:
    joined = " ".join(str(c or "").lower() for c in row)
    return any(k in joined for k in ("total", "jumlah", "laba bersih", "saldo akhir", "saldo awal"))


def generate_excel_report(
    headers: list[str],
    rows: list[list[Any]],
    sheet_title: str = "Report",
    *,
    subtitle: str = "",
    landscape: bool = False,
    generated_by: str = "",
) -> bytes:
    """Build audit-ready .xlsx with letterhead, freeze panes, currency, signatures."""
    branding = get_report_branding()
    primary = branding.primary_color
    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_title or "Report")[:31]

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor=primary)
    total_fill = PatternFill("solid", fgColor="DCE8F8")
    currency_format = '"Rp"#,##0.00'
    money_cols = {i + 1 for i, h in enumerate(headers) if _is_money_header(h)}
    span = max(len(headers), 3)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    ws.cell(row=1, column=1, value=branding.org_name).font = Font(bold=True, size=14, color=primary)

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
    ws.cell(row=2, column=1, value=branding.org_legal_name).font = Font(size=9, color="444444")

    addr = branding.address_block.replace("\n", " | ") if branding.address_block else ""
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=span)
    ws.cell(row=3, column=1, value=addr).font = Font(size=8, color="555555")

    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=span)
    title_cell = ws.cell(row=5, column=1, value=sheet_title)
    title_cell.font = Font(bold=True, size=12)
    title_cell.alignment = Alignment(horizontal="center")

    row_cursor = 6
    if subtitle:
        ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=span)
        sc = ws.cell(row=row_cursor, column=1, value=subtitle)
        sc.font = Font(size=9, color="444444")
        sc.alignment = Alignment(horizontal="center")
        row_cursor += 1

    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    meta = f"Dicetak: {now}"
    if generated_by:
        meta += f" · oleh {generated_by}"
    ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=span)
    mc = ws.cell(row=row_cursor, column=1, value=meta)
    mc.font = Font(size=8, color="666666")
    mc.alignment = Alignment(horizontal="center")
    row_cursor += 2

    header_row = row_cursor
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for r_idx, row in enumerate(rows, start=header_row + 1):
        is_total = _row_is_total(row)
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = border
            if is_total:
                cell.fill = total_fill
                cell.font = Font(bold=True)
            if c_idx in money_cols and isinstance(value, (int, float, Decimal)):
                cell.number_format = currency_format
                cell.alignment = Alignment(horizontal="right")
            elif c_idx in money_cols:
                cell.alignment = Alignment(horizontal="right")

    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        max_len = len(str(headers[col - 1])) if headers else 10
        for row in rows:
            if col - 1 < len(row) and row[col - 1] is not None:
                max_len = max(max_len, min(len(str(row[col - 1])), 40))
        ws.column_dimensions[letter].width = max(max_len + 4, 12)

    ws.freeze_panes = f"A{header_row + 1}"
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.75

    sig_row = header_row + len(rows) + 3
    third = max(span // 3, 1)
    blocks = [
        (1, branding.signatory_left_title, branding.signatory_left_name),
        (1 + third, branding.signatory_mid_title, branding.signatory_mid_name),
        (1 + 2 * third, branding.signatory_right_title, branding.signatory_right_name),
    ]
    for start_col, role, name in blocks:
        end_col = min(start_col + third - 1, span)
        if end_col < start_col:
            end_col = start_col
        ws.merge_cells(start_row=sig_row, start_column=start_col, end_row=sig_row, end_column=end_col)
        role_cell = ws.cell(row=sig_row, column=start_col, value=role)
        role_cell.font = Font(bold=True, size=9)
        role_cell.alignment = Alignment(horizontal="center")

        name_row = sig_row + 4
        ws.merge_cells(start_row=name_row, start_column=start_col, end_row=name_row, end_column=end_col)
        name_cell = ws.cell(
            row=name_row,
            column=start_col,
            value=name or "(........................)",
        )
        name_cell.font = Font(size=9)
        name_cell.alignment = Alignment(horizontal="center")
        name_cell.border = Border(top=Side(style="thin", color="333333"))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
