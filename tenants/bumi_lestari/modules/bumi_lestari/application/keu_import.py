"""Bounded server-side CSV/XLSX preview and atomic application of validated batches."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile
from datetime import date, datetime
from decimal import Decimal

from pydantic import ValidationError
from sqlalchemy import select

from tenants.bumi_lestari.modules.bumi_lestari.application import keu_services as svc, schemas_keu as sc
from tenants.bumi_lestari.modules.bumi_lestari.infrastructure import models_keu as m

HEADERS = {
    "order": ["sumber_ref", "item_ref", "nomor", "tanggal", "produk_id", "nama_snapshot", "qty", "harga_satuan", "pelanggan_id", "varian_snapshot"],
    "biaya": ["sumber_ref", "tanggal", "akun_id", "kategori_id", "jumlah", "keterangan"],
    "settlement": ["sumber_ref", "tanggal_cair", "bruto", "potongan", "penyesuaian", "neto"],
}
PRODUCT_HEADERS = ["sku", "sku_induk", "nama_asli", "gambar_url", "varian_list", "harga_jual"]
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000


def template(kind):
    if kind not in HEADERS:
        svc.bad("Jenis impor tidak dikenal")
    output = io.StringIO()
    csv.writer(output).writerow(HEADERS[kind] + (PRODUCT_HEADERS if kind == "order" else []))
    return output.getvalue()


def parse_file(content, filename):
    if not content or len(content) > MAX_BYTES:
        svc.bad("File harus berisi data dan maksimal 5 MB")
    if filename.lower().endswith(".csv"):
        try:
            raw = content.decode("utf-8-sig")
            dialect = csv.Sniffer().sniff(raw.splitlines()[0], delimiters=",;")
            rows = csv.reader(io.StringIO(raw), dialect)
            return read_rows(rows)
        except (UnicodeDecodeError, csv.Error):
            svc.bad("CSV harus UTF-8 dan memakai delimiter koma atau titik koma")
    if not filename.lower().endswith(".xlsx"):
        svc.bad("Gunakan CSV atau XLSX")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 50 * 1024 * 1024:
                svc.bad("Isi Excel terlalu besar")
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            def values():
                for row in workbook.worksheets[0].iter_rows():
                    if any(cell.data_type == "f" for cell in row):
                        svc.bad("Formula Excel tidak diterima; gunakan nilai hasilnya")
                    yield [cell.value for cell in row]
            return read_rows(values())
        finally:
            workbook.close()
    except (zipfile.BadZipFile, ValueError, KeyError, OSError):
        svc.bad("File XLSX tidak dapat dibaca")


def read_rows(rows):
    iterator = iter(rows)
    header = [str(value or "").strip() for value in next(iterator, [])]
    if not header or len(header) > 32 or len(header) != len(set(header)) or any(not key for key in header):
        svc.bad("Header kosong, duplikat, atau terlalu banyak kolom")
    result = []
    for number, values in enumerate(iterator, 2):
        if number > MAX_ROWS + 1:
            svc.bad("Maksimal 10.000 baris per file")
        if all(value in (None, "") for value in values):
            continue
        if len(values) > len(header) and any(v not in (None, "") for v in values[len(header):]):
            svc.bad(f"Baris {number} memiliki kolom tambahan")
        row = {}
        for key, value in zip(header, list(values) + [None] * len(header)):
            if isinstance(value, (date, datetime)):
                value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
            row[key] = str(value).strip() if value is not None else ""
        result.append((number, row))
    if not result:
        svc.bad("File tidak memiliki baris data")
    return header, result


def validated(kind, channel_id, row):
    if kind == "order":
        q = row.get("qty", "")
        if not q.isdigit():
            raise ValueError("qty harus bilangan bulat positif")
        metadata = None
        if row.get("sku"):
            variants = json.loads(row["varian_list"]) if row.get("varian_list") else ([{"kategori": "Varian", "nilai": row["varian_snapshot"]}] if row.get("varian_snapshot") else [])
            metadata = sc.ProdukSumberIn(sku=row["sku"], sku_induk=row.get("sku_induk") or None,
                nama_asli=row.get("nama_asli") or row["nama_snapshot"], gambar_url=row.get("gambar_url") or "",
                varian_list=variants, harga_jual=row.get("harga_jual") or row["harga_satuan"])
        elif any(row.get(key) for key in PRODUCT_HEADERS if key != "sku"):
            raise ValueError("Metadata produk memerlukan SKU varian")
        line = sc.ItemIn(produk_sumber=metadata, sumber_ref=row["item_ref"], produk_id=row.get("produk_id") or None,
                         nama_snapshot=row["nama_snapshot"], varian_snapshot=row.get("varian_snapshot", ""),
                         qty=int(q), harga_satuan=row["harga_satuan"],
                         subtotal_sumber=Decimal(row["harga_satuan"]) * int(q))
        return sc.PesananIn(saluran_id=channel_id, sumber_ref=row["sumber_ref"], nomor=row["nomor"],
                            tanggal=row["tanggal"], pelanggan_id=row.get("pelanggan_id") or None,
                            status_sumber="manual", total_sumber=line.subtotal_sumber, items=[line])
    if kind == "biaya":
        return sc.BiayaOperasionalIn(saluran_id=channel_id, **row)
    return sc.SettlementIn(saluran_id=channel_id, **{**row, "penyesuaian": row.get("penyesuaian") or "0"})


async def detail(session, key):
    batch = await svc.get(session, m.KeuImpor, key)
    entries = (await session.execute(select(m.KeuMasukan).where(m.KeuMasukan.impor_id == key).order_by(m.KeuMasukan.nomor_baris))).scalars().all()
    return {**svc.record(batch), "rows": [svc.record(row) for row in entries]}


async def preview(session, user, channel_id, kind, filename, content):
    await svc.channel(session, channel_id, manual=kind in {"order", "biaya"})
    if kind not in HEADERS:
        svc.bad("Jenis impor tidak dikenal")
    digest = hashlib.sha256(content).hexdigest()
    await session.execute(select(m.KeuSaluran).where(m.KeuSaluran.id == channel_id).with_for_update())
    previous = (await session.execute(select(m.KeuImpor).where(m.KeuImpor.saluran_id == channel_id, m.KeuImpor.jenis == kind, m.KeuImpor.file_sha256 == digest, m.KeuImpor.versi_format == 1))).scalar_one_or_none()
    if previous:
        return await detail(session, previous.id)
    header, rows = parse_file(content, filename)
    required = set(HEADERS[kind])
    allowed = required | (set(PRODUCT_HEADERS) if kind == "order" else set())
    if not required.issubset(header) or not set(header).issubset(allowed):
        svc.bad("Kolom file harus sesuai template")
    batch = m.KeuImpor(saluran_id=channel_id, jenis=kind, nama_file=filename[:255], file_sha256=digest,
                       versi_format=1, pemetaan={key: key for key in header}, dibuat_oleh=user.id)
    session.add(batch)
    await session.flush()
    failed = False
    seen = set()
    for number, row in rows:
        errors = []
        try:
            validated(kind, channel_id, row)
            identity = (row.get("sumber_ref"), row.get("item_ref"))
            if identity in seen:
                raise ValueError("Referensi baris duplikat")
            seen.add(identity)
        except (ValidationError, ValueError, KeyError, ArithmeticError) as exc:
            failed = True
            errors = [{"pesan": "Baris tidak valid; periksa kolom template dan nominal" if isinstance(exc, ValidationError) else str(exc)[:255]}]
        source_ref = row.get("sumber_ref") or f"file:{digest}:{number}"
        if len(source_ref) > 255:
            source_ref = f"file:{digest}:{number}"
        envelope = sc.MasukanIn(saluran_id=channel_id, entitas=kind, sumber_ref=source_ref,
                                impor_id=batch.id, nomor_baris=number, payload=row)
        session.add(m.KeuMasukan(**envelope.model_dump(), revisi_sha256=sc.checksum_revisi(envelope),
                                 status="gagal" if errors else "menunggu", kesalahan=errors))
    batch.status = "draf" if failed else "valid"
    await session.flush()
    # Cross-row references/headers are checked before the batch is applied.
    if not failed:
        try:
            grouped(kind, channel_id, rows)
        except ValueError:
            batch.status = "draf"
            for entry in (await session.execute(select(m.KeuMasukan).where(m.KeuMasukan.impor_id == batch.id))).scalars():
                entry.status, entry.kesalahan = "gagal", [{"pesan": "Header pesanan dengan referensi sama tidak konsisten"}]
    await svc.audit(session, user, batch, "pratinjau")
    await session.flush()
    return await detail(session, batch.id)


def grouped(kind, channel_id, rows):
    output = {}
    if kind != "order":
        return [validated(kind, channel_id, row) for _, row in rows]
    for _, row in rows:
        payload = validated(kind, channel_id, row)
        old = output.get(payload.sumber_ref)
        if old:
            if (old.nomor, old.tanggal, old.pelanggan_id) != (payload.nomor, payload.tanggal, payload.pelanggan_id):
                raise ValueError("Header tidak konsisten")
            old.items.extend(payload.items)
            old.total_sumber += payload.total_sumber
        else:
            output[payload.sumber_ref] = payload
    return [sc.PesananIn.model_validate(value.model_dump()) for value in output.values()]


async def apply(session, user, key):
    batch = await svc.get(session, m.KeuImpor, key, lock=True)
    if batch.status == "diterapkan":
        return await detail(session, key)
    if batch.status != "valid":
        svc.bad("Perbaiki kesalahan impor sebelum menerapkan", 409)
    entries = (await session.execute(select(m.KeuMasukan).where(m.KeuMasukan.impor_id == key).order_by(m.KeuMasukan.nomor_baris).with_for_update())).scalars().all()
    payloads = grouped(batch.jenis, batch.saluran_id, [(row.nomor_baris, row.payload) for row in entries])
    # Outer request transaction rolls back every row if any reference/rule fails.
    for payload in payloads:
        if batch.jenis == "order":
            await svc.create_order(session, user, payload)
        elif batch.jenis == "biaya":
            await svc.transaction(session, user, payload)
        else:
            await svc.create_settlement(session, user, payload)
    for entry in entries:
        entry.status = "terproses"
    batch.status = "diterapkan"
    await session.flush()
    await svc.audit(session, user, batch, "terapkan")
    return await detail(session, key)
