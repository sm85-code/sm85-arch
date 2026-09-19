"""Load account taxonomy from data/coa_taxonomy.xlsx — single source of slugs."""
from __future__ import annotations

from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

from shared.config import ROOT_DIR

TAXONOMY_PATH = ROOT_DIR / "data" / "coa_taxonomy.xlsx"
COA_PATH = ROOT_DIR / "data" / "coa_code.xlsx"

SUB_KAS_BANK = "kas_bank"
SUB_SALDO_LABA = "saldo_laba"
SUB_IKHTISAR_LR = "ikhtisar_laba_rugi"
SUB_MODAL_DESA = "modal_desa"
SUB_MODAL_MASYARAKAT = "modal_masyarakat"
SUB_BAGI_HASIL_DESA = "bagi_hasil_desa"
SUB_BAGI_HASIL_MASYARAKAT = "bagi_hasil_masyarakat"
SUB_LABA_DICADANGKAN = "laba_dicadangkan"


@lru_cache(maxsize=1)
def load_taxonomy_rows(path: Path | None = None) -> tuple[dict, ...]:
    workbook = path or TAXONOMY_PATH
    if not workbook.is_file():
        return tuple()
    wb = load_workbook(workbook, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows: list[dict] = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if not raw or not raw[0] or not raw[1]:
            continue
        category = str(raw[0]).strip().lower()
        subcategory = str(raw[1]).strip().lower()
        rows.append(
            {
                "category": category,
                "subcategory": subcategory,
                "label_category": str(raw[2] or category).strip(),
                "label_subcategory": str(raw[3] or subcategory).strip(),
                "normal_balance": str(raw[4] or "").strip().lower(),
                "is_system": bool(raw[5]) if len(raw) > 5 else False,
            }
        )
    wb.close()
    return tuple(rows)


def categories_map() -> dict[str, set[str]]:
    out: dict[str, set[str]] = {}
    for row in load_taxonomy_rows():
        out.setdefault(row["category"], set()).add(row["subcategory"])
    return out


def valid_pair(category: str, subcategory: str) -> bool:
    return subcategory in categories_map().get((category or "").strip().lower(), set())
