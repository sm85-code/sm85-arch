"""SIABUMDES B2 — dynamic COA import groups + clone-on-create template.

Unit tests (no Postgres required). DATABASE_URL placeholder for model import.
"""
from __future__ import annotations

import os
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

os.environ.setdefault("DATABASE_URL", "postgresql://user:pass@127.0.0.1:5432/siabumdes_test")

import pytest

from modules.siabumdes.application.coa_template import (
    clone_coa_and_tx_types,
    coa_group_codes,
    find_coa_template_unit,
)
from modules.siabumdes.infrastructure.models import Account, TransactionType, UnitUsaha


def _scalar_result(value):
    result = MagicMock()
    result.scalar_one_or_none.return_value = value
    result.scalars.return_value = value if isinstance(value, list) else ([value] if value is not None else [])
    return result


@pytest.mark.asyncio
async def test_coa_group_codes_includes_bumdes_and_db_units():
    session = MagicMock()
    session.execute = AsyncMock(
        return_value=MagicMock(scalars=MagicMock(return_value=["UU01", "UU99", "UU05"]))
    )
    groups = await coa_group_codes(session)
    assert groups == {"BUMDES", "UU01", "UU99", "UU05"}


@pytest.mark.asyncio
async def test_coa_group_codes_bumdes_only_when_no_units():
    session = MagicMock()
    session.execute = AsyncMock(return_value=MagicMock(scalars=MagicMock(return_value=[])))
    groups = await coa_group_codes(session)
    assert groups == {"BUMDES"}


@pytest.mark.asyncio
async def test_find_template_prefers_uu05_for_perdagangan():
    uu06 = UnitUsaha(id="id6", code="UU06", name="Online", business_type="perdagangan", active=True)
    uu05 = UnitUsaha(id="id5", code="UU05", name="Offline", business_type="perdagangan", active=True)
    session = MagicMock()

    # first execute: candidate list; then scalar counts per candidate (UU05 first)
    candidates_result = MagicMock()
    candidates_result.scalars.return_value = [uu06, uu05]

    count_results = [
        MagicMock(__bool__=lambda self: True),  # unused pattern
    ]

    async def execute_side_effect(stmt):
        # distinguish by looking at compiled-ish string if needed; use call order
        return candidates_result

    call_n = {"n": 0}

    async def scalar_side_effect(stmt):
        # ordered: UU05 then UU06 (preferred order)
        call_n["n"] += 1
        if call_n["n"] == 1:
            return 17  # UU05 has accounts
        return 17

    session.execute = AsyncMock(side_effect=execute_side_effect)
    session.scalar = AsyncMock(side_effect=scalar_side_effect)

    found = await find_coa_template_unit(session, "perdagangan")
    assert found is not None
    assert found.code == "UU05"


@pytest.mark.asyncio
async def test_find_template_explicit_clone_from():
    src = UnitUsaha(id="x", code="UU07", name="Custom", business_type="jasa", active=True)
    session = MagicMock()
    session.execute = AsyncMock(return_value=_scalar_result(src))
    found = await find_coa_template_unit(session, "jasa", clone_from_code="uu07")
    assert found is src


@pytest.mark.asyncio
async def test_find_template_skips_units_without_accounts():
    empty = UnitUsaha(id="e", code="UU01", name="Empty", business_type="jasa", active=True)
    filled = UnitUsaha(id="f", code="UU03", name="Filled", business_type="jasa", active=True)
    session = MagicMock()
    session.execute = AsyncMock(
        return_value=MagicMock(scalars=MagicMock(return_value=[empty, filled]))
    )
    counts = {"UU01": 0, "UU03": 14}
    # find iterates preferred UU01,UU02,UU03,UU04 then remaining
    order = []

    async def scalar_side_effect(stmt):
        # We can't easily read code from stmt; return based on call order matching
        # preferred: UU01 (empty), UU03 (filled) — UU02/UU04 not in candidates
        order.append(1)
        if len(order) == 1:
            return 0  # UU01
        return 14  # UU03

    session.scalar = AsyncMock(side_effect=scalar_side_effect)
    found = await find_coa_template_unit(session, "jasa")
    assert found is not None
    assert found.code == "UU03"


@pytest.mark.asyncio
async def test_clone_coa_and_tx_types_copies_and_skips_existing():
    source = UnitUsaha(id="src", code="UU05", name="Src", business_type="perdagangan", active=True)
    target = UnitUsaha(id="dst", code="UU07", name="Dst", business_type="perdagangan", active=True)

    src_acc = Account(
        id="a1",
        code="1.1.01.01",
        name="Kas",
        category="aset",
        subcategory="kas_bank",
        normal_balance="debit",
        group_code="UU05",
        unit_usaha_id="src",
        active=True,
    )
    already = Account(
        id="a2",
        code="1.1.01.01",
        name="Kas existing",
        category="aset",
        subcategory="kas_bank",
        normal_balance="debit",
        group_code="UU07",
        unit_usaha_id="dst",
        active=True,
    )
    extra = Account(
        id="a3",
        code="1.1.02.01",
        name="Bank",
        category="aset",
        subcategory="kas_bank",
        normal_balance="debit",
        group_code="UU05",
        unit_usaha_id="src",
        active=True,
    )
    src_tt = TransactionType(
        id="t1",
        code="penjualan",
        name="Penjualan",
        debit="1.1.01.01",
        credit="4.1.01.01",
        group_code="UU05",
        unit_codes=["UU05"],
    )

    session = MagicMock()
    session.add = MagicMock()
    session.flush = AsyncMock()

    # execute order in clone_coa_and_tx_types:
    # 1) src accounts, 2) existing target account codes, 3) src types, 4) existing target type codes
    results = [
        MagicMock(scalars=MagicMock(return_value=[src_acc, extra])),
        MagicMock(scalars=MagicMock(return_value=["1.1.01.01"])),  # existing codes
        MagicMock(scalars=MagicMock(return_value=[src_tt])),
        MagicMock(scalars=MagicMock(return_value=[])),
    ]
    session.execute = AsyncMock(side_effect=results)

    stats = await clone_coa_and_tx_types(session, source=source, target=target)
    assert stats["source_unit"] == "UU05"
    assert stats["accounts_cloned"] == 1  # only Bank; Kas already exists
    assert stats["transaction_types_cloned"] == 1
    assert session.add.call_count == 2
    added_acc = session.add.call_args_list[0].args[0]
    assert isinstance(added_acc, Account)
    assert added_acc.code == "1.1.02.01"
    assert added_acc.group_code == "UU07"
    assert added_acc.unit_usaha_id == "dst"
    added_tt = session.add.call_args_list[1].args[0]
    assert isinstance(added_tt, TransactionType)
    assert added_tt.group_code == "UU07"
    assert added_tt.unit_codes == ["UU07"]


@pytest.mark.asyncio
async def test_import_accounts_accepts_dynamic_unit_sheet(monkeypatch):
    """Router-level: valid_groups comes from DB, so UU99 sheet is imported."""
    from io import BytesIO

    from openpyxl import Workbook

    from modules.siabumdes.adapters.api.v1 import master_data_router as mdr

    wb = Workbook()
    # default sheet -> rename to UU99
    ws = wb.active
    ws.title = "UU99"
    ws.append(["code", "name", "category", "subcategory", "normal_balance"])
    ws.append(["1.1.01.01", "Kas", "aset", "kas_bank", "debit"])
    # unknown sheet must be ignored
    ws2 = wb.create_sheet("NOT_A_UNIT")
    ws2.append(["code", "name", "category", "subcategory", "normal_balance"])
    ws2.append(["9.9.99.99", "X", "aset", "kas_bank", "debit"])
    buf = BytesIO()
    wb.save(buf)
    content = buf.getvalue()

    unit = UnitUsaha(id="u99", code="UU99", name="New", business_type="jasa", active=True)
    session = MagicMock()
    session.add = MagicMock()
    session.flush = AsyncMock()

    async def fake_groups(_session):
        return {"BUMDES", "UU99"}

    monkeypatch.setattr(mdr, "coa_group_codes", fake_groups)
    monkeypatch.setattr(mdr, "_validate_category_pair", lambda *a, **k: None)
    monkeypatch.setattr(mdr, "record_audit", AsyncMock())

    # execute: unit lookup for UU99, then exists check for account
    session.execute = AsyncMock(
        side_effect=[
            MagicMock(scalar_one_or_none=MagicMock(return_value=unit)),  # unit for sheet
            MagicMock(scalar_one_or_none=MagicMock(return_value=None)),  # account not exists
        ]
    )

    file = MagicMock()
    file.read = AsyncMock(return_value=content)
    actor = SimpleNamespace(id="admin", role="admin")
    request = MagicMock()
    request.client = SimpleNamespace(host="127.0.0.1")

    result = await mdr.import_accounts(request=request, file=file, actor=actor, session=session)
    assert result["inserted"] == 1
    assert session.add.call_count == 1
    added = session.add.call_args.args[0]
    assert added.group_code == "UU99"
    assert added.unit_usaha_id == "u99"


@pytest.mark.asyncio
async def test_import_accounts_skips_unknown_sheet_not_in_db(monkeypatch):
    from io import BytesIO

    from openpyxl import Workbook

    from modules.siabumdes.adapters.api.v1 import master_data_router as mdr

    wb = Workbook()
    ws = wb.active
    ws.title = "UU88"
    ws.append(["code", "name", "category", "subcategory", "normal_balance"])
    ws.append(["1.1.01.01", "Kas", "aset", "kas_bank", "debit"])
    buf = BytesIO()
    wb.save(buf)

    session = MagicMock()
    session.add = MagicMock()
    monkeypatch.setattr(mdr, "coa_group_codes", AsyncMock(return_value={"BUMDES", "UU01"}))
    monkeypatch.setattr(mdr, "record_audit", AsyncMock())

    file = MagicMock()
    file.read = AsyncMock(return_value=buf.getvalue())
    actor = SimpleNamespace(id="admin", role="admin")
    request = MagicMock()
    request.client = SimpleNamespace(host="127.0.0.1")

    result = await mdr.import_accounts(request=request, file=file, actor=actor, session=session)
    assert result["inserted"] == 0
    session.add.assert_not_called()


@pytest.mark.asyncio
async def test_create_unit_clones_coa_by_default(monkeypatch):
    from modules.siabumdes.adapters.api.v1 import master_data_router as mdr

    session = MagicMock()
    session.add = MagicMock()
    session.flush = AsyncMock()
    # exists check -> None
    session.execute = AsyncMock(return_value=MagicMock(scalar_one_or_none=MagicMock(return_value=None)))

    src = UnitUsaha(id="src", code="UU01", name="Template", business_type="jasa", active=True)
    monkeypatch.setattr(mdr, "find_coa_template_unit", AsyncMock(return_value=src))
    monkeypatch.setattr(
        mdr,
        "clone_coa_and_tx_types",
        AsyncMock(return_value={"source_unit": "UU01", "accounts_cloned": 14, "transaction_types_cloned": 3}),
    )
    monkeypatch.setattr(mdr, "record_audit", AsyncMock())

    payload = mdr.UnitIn(code="uu10", name="Unit Baru", business_type="jasa")
    request = MagicMock()
    request.client = SimpleNamespace(host="127.0.0.1")
    actor = SimpleNamespace(id="admin", role="admin")

    out = await mdr.create_unit(payload=payload, request=request, actor=actor, session=session)
    assert out["code"] == "UU10"
    assert out["coa_template"]["cloned"] is True
    assert out["coa_template"]["accounts_cloned"] == 14
    mdr.clone_coa_and_tx_types.assert_awaited_once()


@pytest.mark.asyncio
async def test_create_unit_clone_coa_false_skips(monkeypatch):
    from modules.siabumdes.adapters.api.v1 import master_data_router as mdr

    session = MagicMock()
    session.add = MagicMock()
    session.flush = AsyncMock()
    session.execute = AsyncMock(return_value=MagicMock(scalar_one_or_none=MagicMock(return_value=None)))
    find = AsyncMock()
    monkeypatch.setattr(mdr, "find_coa_template_unit", find)
    monkeypatch.setattr(mdr, "record_audit", AsyncMock())

    payload = mdr.UnitIn(code="UU11", name="Empty", business_type="jasa", clone_coa=False)
    request = MagicMock()
    request.client = SimpleNamespace(host="127.0.0.1")
    actor = SimpleNamespace(id="admin", role="admin")

    out = await mdr.create_unit(payload=payload, request=request, actor=actor, session=session)
    assert out["coa_template"] == {"cloned": False}
    find.assert_not_awaited()
