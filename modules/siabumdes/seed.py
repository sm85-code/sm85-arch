"""Startup seed: taxonomy + COA + default users. Idempotent, no overwrite of live data."""
from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert as pg_insert

from modules.siabumdes.identity.infrastructure.models import SystemControl, User
from modules.siabumdes.infrastructure.models import (
    Account,
    AccountCategory,
    AccountSubcategory,
    UnitUsaha,
)
from modules.siabumdes.coa_taxonomy import COA_PATH, TAXONOMY_PATH, load_taxonomy_rows, valid_pair
from shared.database import SessionLocal
from shared.security import hash_password

logger = logging.getLogger("sm85.seed")

UNIT_CATALOG = {
    "BUMDES": ("BUMDes (Pusat)", "Kantor pusat BUMDes"),
    "UU01": ("Pembibitan Domba Garut", "Unit usaha UU01"),
    "UU02": ("Peternakan Ikan Air Tawar Sistem Bioflok", "Unit usaha UU02"),
    "UU03": ("Sewa Kendaraan Angkutan", "Unit usaha UU03"),
    "UU04": ("Karya Raharja Sinergi Digital (KRSD)", "Unit usaha UU04"),
    "UU05": ("Toko Offline BUMDES", "Unit usaha UU05 — inventori"),
    "UU06": ("Toko Online BUMDES", "Unit usaha UU06"),
}

DEFAULT_USERS = [
    ("admin", "admin123@", "Administrator", "admin", None),
    ("direktur", "direktur123@", "Direktur", "direktur", None),
    ("bendahara", "bendahara123@", "Bendahara", "bendahara", None),
    ("penasihat", "penasihat123@", "Penasihat", "penasihat", None),
    ("pengawas", "pengawas123@", "Pengawas", "pengawas", None),
    ("pengelola1", "unit1.iyes", "Pengelola UU01", "pengelola", "UU01"),
    ("pengelola2", "unit2.iyes", "Pengelola UU02", "pengelola", "UU02"),
    ("pengelola3", "unit3.iyes", "Pengelola UU03", "pengelola", "UU03"),
    ("pengelola4", "unit4.iyes", "Pengelola UU04", "pengelola", "UU04"),
    ("pengelola5", "unit5.iyes", "Pengelola UU05", "pengelola", "UU05"),
    ("pengelola6", "unit6.iyes", "Pengelola UU06", "pengelola", "UU06"),
]


def _iter_coa_rows(path: Path):
    wb = load_workbook(path, data_only=True, read_only=True)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(c or "").strip().lower() for c in rows[0]]
        if header[:5] != ["code", "name", "category", "subcategory", "normal_balance"]:
            continue
        for raw in rows[1:]:
            if not raw or not raw[0]:
                continue
            code = str(raw[0]).strip()
            name = str(raw[1] or "").strip()
            category = str(raw[2] or "").strip().lower()
            subcategory = str(raw[3] or "").strip().lower()
            normal_balance = str(raw[4] or "").strip().lower()
            group = str(raw[5] or sheet).strip().upper() if len(raw) > 5 else sheet.upper()
            if not name or normal_balance not in {"debit", "kredit"}:
                continue
            if not valid_pair(category, subcategory):
                logger.warning("skip COA %s %s: unknown %s/%s", group, code, category, subcategory)
                continue
            yield {
                "code": code,
                "name": name,
                "category": category,
                "subcategory": subcategory,
                "normal_balance": normal_balance,
                "group_code": group,
            }
    wb.close()


async def seed_if_needed() -> None:
    async with SessionLocal() as session:
        try:
            units = await _seed_units(session)
            await _seed_taxonomy(session)
            await _seed_coa(session, units)
            await _seed_users(session, units)
            existing = await session.get(SystemControl, "default")
            if not existing:
                session.add(SystemControl(id="default", recording_locked=False))
            await session.commit()
            logger.info("seed completed (idempotent) from %s + %s", TAXONOMY_PATH.name, COA_PATH.name)
        except Exception:
            await session.rollback()
            logger.exception("seed failed")
            raise


async def _seed_units(session) -> dict[str, UnitUsaha]:
    existing = {u.code: u for u in (await session.execute(select(UnitUsaha))).scalars()}
    created = 0
    for code, (name, desc) in UNIT_CATALOG.items():
        if code == "BUMDES":
            continue
        if code in existing:
            continue
        row = UnitUsaha(code=code, name=name, description=desc, active=True)
        session.add(row)
        existing[code] = row
        created += 1
    await session.flush()
    logger.info("unit seed created=%s existing=%s", created, len(existing))
    return existing


async def _seed_taxonomy(session) -> None:
    rows = load_taxonomy_rows()
    if not rows:
        logger.warning("taxonomy workbook missing at %s", TAXONOMY_PATH)
        return
    cats = [
        {"slug": item["category"], "label": item["label_category"], "normal_balance": item["normal_balance"] or "debit"}
        for item in rows
    ]
    seen: set[str] = set()
    cat_payload = []
    for row in cats:
        if row["slug"] in seen:
            continue
        seen.add(row["slug"])
        cat_payload.append(row)
    if cat_payload:
        await session.execute(
            pg_insert(AccountCategory).values(cat_payload).on_conflict_do_nothing(index_elements=["slug"])
        )
    sub_payload = []
    seen_sub: set[str] = set()
    for item in rows:
        slug = item["subcategory"]
        if slug in seen_sub:
            continue
        seen_sub.add(slug)
        sub_payload.append(
            {
                "slug": slug,
                "category_slug": item["category"],
                "label": item["label_subcategory"],
                "is_system": item["is_system"],
            }
        )
    if sub_payload:
        await session.execute(
            pg_insert(AccountSubcategory).values(sub_payload).on_conflict_do_nothing(index_elements=["slug"])
        )
    await session.flush()
    logger.info("taxonomy upsert categories=%s subcategories=%s", len(cat_payload), len(sub_payload))


async def _seed_coa(session, units: dict[str, UnitUsaha]) -> None:
    if not COA_PATH.is_file():
        logger.warning("COA workbook missing at %s", COA_PATH)
        return
    payload = []
    seen: set[tuple[str, str]] = set()
    for item in _iter_coa_rows(COA_PATH):
        key = (item["code"], item["group_code"])
        if key in seen:
            continue
        seen.add(key)
        unit = units.get(item["group_code"])
        payload.append(
            {
                "code": item["code"],
                "name": item["name"],
                "category": item["category"],
                "subcategory": item["subcategory"],
                "normal_balance": item["normal_balance"],
                "group_code": item["group_code"],
                "unit_usaha_id": unit.id if unit else None,
                "active": True,
            }
        )
    inserted = 0
    if payload:
        result = await session.execute(
            pg_insert(Account).values(payload).on_conflict_do_nothing(constraint="uq_accounts_code_group")
        )
        inserted = result.rowcount or 0
    await session.flush()
    logger.info("coa seed inserted=%s skipped_existing=%s from %s", inserted, max(len(payload) - inserted, 0), COA_PATH.name)


async def _seed_users(session, units: dict[str, UnitUsaha]) -> None:
    """Cuma seed akun default kalau tabel users BENAR-BENAR kosong (instalasi
    baru). Dulu tiap username DEFAULT_USERS dicek satu-satu dan yang belum
    ada langsung dibuat ulang -- efeknya, akun default (mis. "direktur")
    yang sengaja dihapus admin lewat Kelola Pengguna (karena sudah diganti
    akun baru) terus hidup lagi tiap kali aplikasi restart/deploy, walau
    seed_if_needed() dipanggil di lifespan startup setiap boot."""
    existing = {u.username: u for u in (await session.execute(select(User))).scalars()}
    if existing:
        logger.info("user seed skipped: users table already has %s row(s)", len(existing))
        return
    for username, password, name, role, unit_code in DEFAULT_USERS:
        unit = units.get(unit_code) if unit_code else None
        session.add(
            User(
                username=username,
                email=f"{username}@siabumdes.local",
                name=name,
                password_hash=hash_password(password),
                role=role,
                unit_usaha_id=unit.id if unit else None,
                must_change_password=False,
                active=True,
            )
        )
    logger.info("user seed created=%s (fresh install)", len(DEFAULT_USERS))
