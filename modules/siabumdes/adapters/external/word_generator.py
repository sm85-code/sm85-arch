"""Audit-ready .docx report generator (letterhead, table, signatures).

Mirrors the layout of modules.siabumdes.adapters.external.pdf_generator / excel_adapter so PDF,
Excel and Word exports of the same report look consistent.
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

from modules.siabumdes.adapters.external.report_formatting import format_money, is_money_header, looks_numeric, row_is_total
from modules.siabumdes.report_branding import ReportBranding, fetch_logo_bytes, get_report_branding

_is_money_header = is_money_header
_row_is_total = row_is_total


def _rgb(hex6: str) -> RGBColor:
    hex6 = (hex6 or "1C8A8A").lstrip("#")
    return RGBColor(int(hex6[0:2], 16), int(hex6[2:4], 16), int(hex6[4:6], 16))


def _shade_cell(cell, hex6: str) -> None:
    shd = cell._tc.get_or_add_tcPr().makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): hex6})
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_text(cell, text: str, *, bold: bool = False, size: int = 9, color: Optional[RGBColor] = None, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run(text if text is not None else "")
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color


def generate_word_report(
    *,
    title: str,
    subtitle: str = "",
    table_headers: Optional[list[str]] = None,
    table_rows: Optional[list[list[Any]]] = None,
    footer: str = "",
    landscape: bool = False,
    branding: Optional[ReportBranding] = None,
    generated_by: str = "",
) -> bytes:
    """Build an audit-ready .docx: letterhead, title, data table, signatures."""
    branding = branding or get_report_branding()
    primary_hex = branding.primary_color or "1C8A8A"
    primary_rgb = _rgb(primary_hex)

    doc = Document()
    section = doc.sections[0]
    if landscape:
        section.orientation = 1  # WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = section.right_margin = Cm(1.6)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.6)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    logo_bytes = fetch_logo_bytes(branding.logo_url) if branding.logo_url else None
    if logo_bytes:
        try:
            logo_p = doc.add_paragraph()
            logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            logo_p.paragraph_format.space_after = Pt(2)
            logo_p.add_run().add_picture(io.BytesIO(logo_bytes), height=Cm(1.6))
        except Exception:  # noqa: BLE001 -- a broken/unsupported logo must never break the report
            pass

    org_p = doc.add_paragraph()
    org_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    org_run = org_p.add_run(branding.org_name)
    org_run.bold = True
    org_run.font.size = Pt(15)
    org_run.font.color.rgb = primary_rgb
    org_p.paragraph_format.space_after = Pt(1)

    legal_p = doc.add_paragraph()
    legal_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    legal_run = legal_p.add_run(branding.org_legal_name)
    legal_run.font.size = Pt(9)
    legal_run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    legal_p.paragraph_format.space_after = Pt(1)

    if branding.address_block:
        addr_p = doc.add_paragraph()
        addr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        addr_run = addr_p.add_run(branding.address_block.replace("\n", " · "))
        addr_run.font.size = Pt(8)
        addr_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        addr_p.paragraph_format.space_after = Pt(4)

    rule_p = doc.add_paragraph()
    rule_fmt = rule_p.paragraph_format
    rule_fmt.space_after = Pt(6)
    pPr = rule_p._p.get_or_add_pPr()
    pbdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pPr.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "18", qn("w:color"): primary_hex})
    pbdr.append(bottom)
    pPr.append(pbdr)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title.upper())
    title_run.bold = True
    title_run.font.size = Pt(13)
    title_p.paragraph_format.space_after = Pt(2)

    if subtitle:
        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub_p.add_run(subtitle)
        sub_run.font.size = Pt(10)
        sub_run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        sub_p.paragraph_format.space_after = Pt(2)

    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    meta = f"Dicetak: {now}"
    if generated_by:
        meta += f" · oleh {generated_by}"
    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta_p.add_run(meta)
    meta_run.font.size = Pt(8)
    meta_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    meta_p.paragraph_format.space_after = Pt(10)

    headers = table_headers or []
    rows = table_rows or []
    if headers and rows is not None:
        money_cols = {i for i, h in enumerate(headers) if _is_money_header(h)}
        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            align = WD_ALIGN_PARAGRAPH.RIGHT if i in money_cols else WD_ALIGN_PARAGRAPH.LEFT
            _set_cell_text(cell, str(h), bold=True, size=9, color=RGBColor(0xFF, 0xFF, 0xFF), align=align)
            _shade_cell(cell, primary_hex)

        for row in rows:
            is_total = _row_is_total(row)
            cells = table.add_row().cells
            for i, value in enumerate(row):
                is_num = i in money_cols and looks_numeric(value)
                if is_num:
                    text = format_money(value)
                else:
                    text = "" if value is None else str(value)
                align = WD_ALIGN_PARAGRAPH.RIGHT if is_num else WD_ALIGN_PARAGRAPH.LEFT
                _set_cell_text(cells[i], text, bold=is_total, size=9, align=align)
                if is_total:
                    _shade_cell(cells[i], "DCE8F8")

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    sign_table = doc.add_table(rows=2, cols=3)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    roles = [
        (branding.signatory_left_title, branding.signatory_left_name),
        (branding.signatory_mid_title, branding.signatory_mid_name),
        (branding.signatory_right_title, branding.signatory_right_name),
    ]
    for col, (role, name) in enumerate(roles):
        _set_cell_text(sign_table.rows[0].cells[col], role, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(
            sign_table.rows[1].cells[col],
            name or "(........................)",
            size=9,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for row in sign_table.rows:
        for cell in row.cells:
            cell._tc.get_or_add_tcPr()

    if footer:
        foot_p = doc.add_paragraph()
        foot_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        foot_p.paragraph_format.space_before = Pt(14)
        foot_run = foot_p.add_run(footer)
        foot_run.font.size = Pt(7.5)
        foot_run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
