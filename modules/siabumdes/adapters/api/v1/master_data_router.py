"""Master data HTTP API matching frontend-siabumdes paths."""
from __future__ import annotations

from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from modules.siabumdes.adapters.api.deps import get_current_user, require_roles
from modules.siabumdes.adapters.api.scope import (
    MASTER_WRITE_ROLES,
    UNIT_WRITE_ROLES,
    is_pengelola,
    pengelola_unit,
    scoped_unit_id,
)
from modules.siabumdes.adapters.external.excel_adapter import generate_excel_report, parse_excel_rows
from modules.siabumdes.application.coa_template import (
    clone_coa_and_tx_types,
    coa_group_codes,
    find_coa_template_unit,
)
from modules.siabumdes.coa_taxonomy import categories_map, valid_pair
from modules.siabumdes.identity.application.services import record_audit
from modules.siabumdes.identity.infrastructure.models import User
from modules.siabumdes.infrastructure.models import Account, Mitra, TransactionType, UnitUsaha
from shared.database import get_db

router = APIRouter(prefix="/api", tags=["master-data"])


def _client_ip(request: Request) -> str:
    return request.client.host if request.client else "unknown"


def _validate_category_pair(category: str, subcategory: str) -> Optional[str]:
    """Validate (category, subcategory) against data/coa_taxonomy.xlsx -- the
    single source of truth for what pairs are allowed (shared with the
    startup seed via modules.siabumdes.coa_taxonomy). Returns an error
    message, or None if valid.

    A blank subcategory is accepted (kept backward compatible with accounts
    that were created without one) -- only a *non-empty* subcategory must
    match a real pair from the taxonomy.
    """
    if category not in categories_map():
        return f"Kategori '{category}' tidak dikenal (lihat data/coa_taxonomy.xlsx)"
    if subcategory and not valid_pair(category, subcategory):
        return f"Subkategori '{subcategory}' tidak valid untuk kategori '{category}' (lihat data/coa_taxonomy.xlsx)"
    return None


class AccountIn(BaseModel):
    code: str
    name: str
    category: str
    subcategory: str = ""
    normal_balance: str
    group: str = "BUMDES"


BUSINESS_TYPES = {"jasa", "perdagangan", "manufaktur"}


class UnitIn(BaseModel):
    code: str
    name: str
    business_type: str = "jasa"
    description: str = ""
    revenue_scheme: str = ""
    # Optional COA/tx-type seed from an existing same-type unit (BE default on).
    # FE may omit; no contract break. Set false to create an empty unit.
    clone_coa: bool = True
    clone_from: Optional[str] = None


class UnitPatch(BaseModel):
    name: Optional[str] = None
    business_type: Optional[str] = None
    description: Optional[str] = None
    revenue_scheme: Optional[str] = None
    active: Optional[bool] = None


class TxTypeIn(BaseModel):
    code: str
    name: str
    debit: str
    credit: str
    unit_codes: list[str] = Field(default_factory=list)
    group: str = "BUMDES"


class MitraIn(BaseModel):
    name: str
    unit_usaha_id: Optional[str] = None
    phone: str = ""
    note: str = ""


def _acc_out(row: Account) -> dict:
    return {
        "id": row.id,
        "code": row.code,
        "name": row.name,
        "category": row.category,
        "subcategory": row.subcategory,
        "normal_balance": row.normal_balance,
        "group": row.group_code,
        "group_code": row.group_code,
        "unit_usaha_id": row.unit_usaha_id,
        "active": row.active,
    }


def _type_out(row: TransactionType) -> dict:
    return {
        "id": row.id,
        "code": row.code,
        "name": row.name,
        "debit": row.debit,
        "credit": row.credit,
        "group": row.group_code,
        "unit_codes": list(row.unit_codes or []),
    }


async def _group_filter(session: AsyncSession, user: User) -> Optional[str]:
    if is_pengelola(user):
        return (await pengelola_unit(session, user)).code
    return None


@router.get("/unit-usaha")
async def list_units(
    include_inactive: bool = False,
    user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_db),
):
    stmt = select(UnitUsaha).order_by(UnitUsaha.code.asc())
    if is_pengelola(user):
        # Selalu tampilkan unit sendiri biarpun sudah dinonaktifkan -- kalau
        # tidak, pengelola yang unitnya dinonaktifkan kehilangan akses ke
        # halamannya sendiri secara tiba-tiba tanpa penjelasan.
        stmt = stmt.where(UnitUsaha.id == user.unit_usaha_id)
    elif not include_inactive:
        # Dropdown pemilih unit di Transaksi/Laporan/COA/dll tidak boleh
        # menawarkan unit yang sudah dinonaktifkan admin. Halaman Profil Unit
        # Usaha sendiri butuh lihat semua (termasuk nonaktif) supaya masih
        # bisa diaktifkan lagi -- itu lewat include_inactive=true eksplisit.
        stmt = stmt.where(UnitUsaha.active.is_(True))
    rows = (await session.execute(stmt)).scalars()
    return [
        {
            "id": row.id,
            "code": row.code,
            "name": row.name,
            "business_type": row.business_type,
            "description": row.description,
            "revenue_scheme": row.revenue_scheme,
            "active": row.active,
        }
        for row in rows
    ]


@router.post("/unit-usaha")
async def create_unit(
    payload: UnitIn,
    request: Request,
    actor: User = Depends(require_roles(*UNIT_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    if payload.business_type not in BUSINESS_TYPES:
        raise HTTPException(status_code=422, detail="Jenis usaha tidak valid")
    exists = (await session.execute(select(UnitUsaha).where(UnitUsaha.code == payload.code.strip()))).scalar_one_or_none()
    if exists:
        raise HTTPException(status_code=400, detail="Kode unit sudah ada")
    if payload.clone_from and payload.clone_coa:
        src = (
            await session.execute(
                select(UnitUsaha).where(UnitUsaha.code == payload.clone_from.strip().upper())
            )
        ).scalar_one_or_none()
        if not src:
            raise HTTPException(status_code=400, detail=f"Unit template {payload.clone_from} tidak ditemukan")
    row = UnitUsaha(
        code=payload.code.strip().upper(),
        name=payload.name.strip(),
        business_type=payload.business_type,
        description=payload.description,
        revenue_scheme=payload.revenue_scheme,
    )
    session.add(row)
    await session.flush()

    clone_info: dict = {"cloned": False}
    if payload.clone_coa:
        template = await find_coa_template_unit(
            session,
            payload.business_type,
            exclude_unit_id=row.id,
            clone_from_code=payload.clone_from,
        )
        if template:
            stats = await clone_coa_and_tx_types(session, source=template, target=row)
            clone_info = {"cloned": True, **stats}
        elif payload.clone_from:
            raise HTTPException(
                status_code=400,
                detail=f"Unit template {payload.clone_from} tidak ditemukan",
            )

    detail = f"Buat unit {row.code} - {row.name}"
    if clone_info.get("cloned"):
        detail += (
            f" (clone COA dari {clone_info.get('source_unit')}: "
            f"{clone_info.get('accounts_cloned')} akun, "
            f"{clone_info.get('transaction_types_cloned')} jenis)"
        )
    await record_audit(
        session, actor=actor, action="create_unit_usaha", entity="unit_usaha", entity_id=row.id,
        detail=detail, ip=_client_ip(request),
    )
    return {
        "id": row.id,
        "code": row.code,
        "name": row.name,
        "business_type": row.business_type,
        "description": row.description,
        "revenue_scheme": row.revenue_scheme,
        "active": row.active,
        "coa_template": clone_info,
    }


@router.patch("/unit-usaha/{unit_id}")
async def update_unit(
    unit_id: str,
    payload: UnitPatch,
    request: Request,
    actor: User = Depends(require_roles(*UNIT_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    row = await session.get(UnitUsaha, unit_id)
    if not row:
        raise HTTPException(status_code=404, detail="Unit usaha tidak ditemukan")
    if payload.business_type is not None and payload.business_type not in BUSINESS_TYPES:
        raise HTTPException(status_code=422, detail="Jenis usaha tidak valid")
    if payload.name is not None:
        row.name = payload.name.strip()
    if payload.business_type is not None:
        row.business_type = payload.business_type
    if payload.description is not None:
        row.description = payload.description
    if payload.revenue_scheme is not None:
        row.revenue_scheme = payload.revenue_scheme
    if payload.active is not None:
        row.active = payload.active
    await session.flush()
    await record_audit(
        session, actor=actor, action="update_unit_usaha", entity="unit_usaha", entity_id=row.id,
        detail=f"Ubah unit {row.code}", ip=_client_ip(request),
    )
    return {
        "id": row.id,
        "code": row.code,
        "name": row.name,
        "business_type": row.business_type,
        "description": row.description,
        "revenue_scheme": row.revenue_scheme,
        "active": row.active,
    }


@router.get("/accounts")
async def list_accounts(user: User = Depends(get_current_user), session: AsyncSession = Depends(get_db)):
    stmt = select(Account).order_by(Account.group_code.asc(), Account.code.asc())
    forced = await _group_filter(session, user)
    if forced:
        stmt = stmt.where(Account.group_code == forced)
    rows = (await session.execute(stmt)).scalars()
    return [_acc_out(row) for row in rows]


@router.post("/accounts")
async def create_account(
    payload: AccountIn,
    request: Request,
    actor: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    group = (payload.group or "BUMDES").strip().upper()
    exists = (
        await session.execute(select(Account).where(Account.code == payload.code, Account.group_code == group))
    ).scalar_one_or_none()
    if exists:
        raise HTTPException(status_code=400, detail=f"Kode akun {payload.code} sudah ada di kelompok {group}")
    category = payload.category.strip().lower()
    subcategory = (payload.subcategory or "").strip().lower()
    error = _validate_category_pair(category, subcategory)
    if error:
        raise HTTPException(status_code=400, detail=error)
    if payload.normal_balance not in {"debit", "kredit"}:
        raise HTTPException(status_code=400, detail="normal_balance harus debit atau kredit")
    unit = (await session.execute(select(UnitUsaha).where(UnitUsaha.code == group))).scalar_one_or_none()
    row = Account(
        code=payload.code.strip(),
        name=payload.name.strip(),
        category=category,
        subcategory=subcategory,
        normal_balance=payload.normal_balance,
        group_code=group,
        unit_usaha_id=unit.id if unit else None,
    )
    session.add(row)
    await session.flush()
    await record_audit(
        session, actor=actor, action="create_account", entity="accounts", entity_id=row.id,
        detail=f"Buat akun {row.code} - {row.name} ({row.group_code})", ip=_client_ip(request),
    )
    return _acc_out(row)


@router.put("/accounts/{code}")
async def update_account(
    code: str,
    payload: AccountIn,
    request: Request,
    group: str = Query("BUMDES"),
    actor: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    row = (
        await session.execute(select(Account).where(Account.code == code, Account.group_code == group))
    ).scalar_one_or_none()
    if not row:
        raise HTTPException(status_code=404, detail=f"Kode akun {code} tidak ditemukan di kelompok {group}")
    category = payload.category.strip().lower()
    subcategory = (payload.subcategory or "").strip().lower()
    error = _validate_category_pair(category, subcategory)
    if error:
        raise HTTPException(status_code=400, detail=error)
    if payload.normal_balance not in {"debit", "kredit"}:
        raise HTTPException(status_code=400, detail="normal_balance harus debit atau kredit")
    target_group = (payload.group or group).strip().upper()
    row.code = payload.code.strip()
    row.name = payload.name.strip()
    row.category = category
    row.subcategory = subcategory
    row.normal_balance = payload.normal_balance
    row.group_code = target_group
    await session.flush()
    await record_audit(
        session, actor=actor, action="update_account", entity="accounts", entity_id=row.id,
        detail=f"Ubah akun {row.code} ({row.group_code})", ip=_client_ip(request),
    )
    return _acc_out(row)


@router.delete("/accounts/{code}")
async def delete_account(
    code: str,
    request: Request,
    group: str = Query("BUMDES"),
    actor: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    row = (
        await session.execute(select(Account).where(Account.code == code, Account.group_code == group))
    ).scalar_one_or_none()
    if not row:
        raise HTTPException(status_code=404, detail="Kode akun tidak ditemukan")
    await record_audit(
        session, actor=actor, action="delete_account", entity="accounts", entity_id=row.id,
        detail=f"Hapus akun {row.code} ({row.group_code})", ip=_client_ip(request),
    )
    await session.delete(row)
    return {"deleted": 1}


@router.delete("/accounts/reset-all")
async def reset_accounts(
    request: Request,
    confirm: str = "",
    actor: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    if confirm != "YES":
        raise HTTPException(status_code=400, detail="Parameter confirm=YES wajib untuk operasi ini")
    rows = (await session.execute(select(Account))).scalars().all()
    for row in rows:
        await session.delete(row)
    await record_audit(
        session, actor=actor, action="reset_accounts", entity="accounts",
        detail=f"Hapus SEMUA akun ({len(rows)} akun)", ip=_client_ip(request),
    )
    return {"deleted": len(rows)}


@router.get("/transaction-types")
async def list_tx_types(user: User = Depends(get_current_user), session: AsyncSession = Depends(get_db)):
    stmt = select(TransactionType).order_by(TransactionType.group_code.asc(), TransactionType.code.asc())
    forced = await _group_filter(session, user)
    if forced:
        stmt = stmt.where(TransactionType.group_code == forced)
    rows = (await session.execute(stmt)).scalars()
    return [_type_out(row) for row in rows]


@router.post("/transaction-types")
async def create_tx_type(
    payload: TxTypeIn,
    request: Request,
    actor: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    group = (payload.group or "BUMDES").strip().upper()
    exists = (
        await session.execute(
            select(TransactionType).where(TransactionType.code == payload.code, TransactionType.group_code == group)
        )
    ).scalar_one_or_none()
    if exists:
        raise HTTPException(status_code=400, detail="Kode jenis transaksi sudah ada")
    row = TransactionType(
        code=payload.code.strip(),
        name=payload.name.strip(),
        debit=payload.debit,
        credit=payload.credit,
        group_code=group,
        unit_codes=payload.unit_codes or [],
    )
    session.add(row)
    await session.flush()
    await record_audit(
        session, actor=actor, action="create_transaction_type", entity="transaction_types", entity_id=row.id,
        detail=f"Buat jenis transaksi {row.code} - {row.name} ({row.group_code})", ip=_client_ip(request),
    )
    return _type_out(row)


@router.put("/transaction-types/{code}")
async def update_tx_type(
    code: str,
    payload: TxTypeIn,
    request: Request,
    actor: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    group = (payload.group or "BUMDES").strip().upper()
    row = (
        await session.execute(
            select(TransactionType).where(TransactionType.code == code, TransactionType.group_code == group)
        )
    ).scalar_one_or_none()
    if not row:
        row = (await session.execute(select(TransactionType).where(TransactionType.code == code))).scalar_one_or_none()
    if not row:
        raise HTTPException(status_code=404, detail="Jenis transaksi tidak ditemukan")
    if payload.code != code:
        clash = (
            await session.execute(
                select(TransactionType).where(
                    TransactionType.code == payload.code, TransactionType.group_code == group
                )
            )
        ).scalar_one_or_none()
        if clash:
            raise HTTPException(status_code=400, detail="Kode tujuan sudah ada")
        row.code = payload.code
    row.name = payload.name
    row.debit = payload.debit
    row.credit = payload.credit
    row.group_code = group
    row.unit_codes = payload.unit_codes or []
    await session.flush()
    await record_audit(
        session, actor=actor, action="update_transaction_type", entity="transaction_types", entity_id=row.id,
        detail=f"Ubah jenis transaksi {row.code} ({row.group_code})", ip=_client_ip(request),
    )
    return _type_out(row)


@router.delete("/transaction-types/{code}")
async def delete_tx_type(
    code: str,
    request: Request,
    actor: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    row = (await session.execute(select(TransactionType).where(TransactionType.code == code))).scalar_one_or_none()
    if not row:
        raise HTTPException(status_code=404, detail="Jenis transaksi tidak ditemukan")
    await record_audit(
        session, actor=actor, action="delete_transaction_type", entity="transaction_types", entity_id=row.id,
        detail=f"Hapus jenis transaksi {row.code} ({row.group_code})", ip=_client_ip(request),
    )
    await session.delete(row)
    return {"deleted": 1}


@router.get("/mitra")
async def list_mitra(
    unit_usaha_id: Optional[str] = None,
    user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_db),
):
    stmt = select(Mitra)
    if is_pengelola(user):
        stmt = stmt.where(Mitra.unit_usaha_id == user.unit_usaha_id)
    elif unit_usaha_id:
        stmt = stmt.where(Mitra.unit_usaha_id == unit_usaha_id)
    rows = (await session.execute(stmt)).scalars()
    return [
        {"id": row.id, "name": row.name, "unit_usaha_id": row.unit_usaha_id, "phone": row.phone, "note": row.note}
        for row in rows
    ]


@router.post("/mitra")
async def create_mitra(
    payload: MitraIn,
    user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_db),
):
    if user.role in {"pengawas", "penasihat"}:
        raise HTTPException(status_code=403, detail="Akun bersifat read-only")
    unit_id = await scoped_unit_id(session, user, payload.unit_usaha_id)
    row = Mitra(name=payload.name.strip(), unit_usaha_id=unit_id, phone=payload.phone, note=payload.note)
    session.add(row)
    await session.flush()
    return {"id": row.id, "name": row.name, "unit_usaha_id": row.unit_usaha_id, "phone": row.phone, "note": row.note}


@router.delete("/mitra/{mitra_id}")
async def delete_mitra(
    mitra_id: str,
    request: Request,
    actor: User = Depends(require_roles("admin", "direktur", "bendahara")),
    session: AsyncSession = Depends(get_db),
):
    row = await session.get(Mitra, mitra_id)
    if not row:
        return {"deleted": 0}
    await record_audit(
        session, actor=actor, action="delete_mitra", entity="mitra", entity_id=row.id,
        detail=f"Hapus mitra {row.name}", ip=_client_ip(request),
    )
    await session.delete(row)
    return {"deleted": 1}


@router.post("/accounts/import")
async def import_accounts(
    request: Request,
    file: UploadFile = File(...),
    actor: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    content = await file.read()
    from openpyxl import load_workbook
    import io as _io

    try:
        wb = load_workbook(filename=_io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"File tidak valid: {exc}") from exc

    inserted = skipped = 0
    errors: list[str] = []
    # Accept BUMDES + any unit code present in DB (not hardcoded UU01–UU06).
    valid_groups = await coa_group_codes(session)
    for sheet_name in wb.sheetnames:
        if sheet_name not in valid_groups:
            continue
        headers, data = parse_excel_rows(content, sheet_name=sheet_name)
        header = [h.strip().lower() for h in headers]
        if header[:5] != ["code", "name", "category", "subcategory", "normal_balance"]:
            errors.append(f"Sheet '{sheet_name}': header tidak sesuai")
            continue
        unit = (await session.execute(select(UnitUsaha).where(UnitUsaha.code == sheet_name))).scalar_one_or_none()
        for idx, values in enumerate(data, start=2):
            code = str(values[0] or "").strip()
            name = str(values[1] or "").strip()
            category = str(values[2] or "").strip().lower()
            subcategory = str(values[3] or "").strip().lower()
            normal_balance = str(values[4] or "").strip().lower()
            if not (code and name and category and normal_balance):
                skipped += 1
                continue
            pair_error = _validate_category_pair(category, subcategory)
            if pair_error or normal_balance not in {"debit", "kredit"}:
                skipped += 1
                errors.append(f"Sheet '{sheet_name}' baris {idx} ({code}): {pair_error or 'normal_balance harus debit atau kredit'}")
                continue
            exists = (
                await session.execute(select(Account).where(Account.code == code, Account.group_code == sheet_name))
            ).scalar_one_or_none()
            if exists:
                skipped += 1
                continue
            session.add(
                Account(
                    code=code,
                    name=name,
                    category=category,
                    subcategory=subcategory,
                    normal_balance=normal_balance,
                    group_code=sheet_name,
                    unit_usaha_id=unit.id if unit else None,
                )
            )
            inserted += 1
    await record_audit(
        session, actor=actor, action="import_accounts", entity="accounts",
        detail=f"Import {inserted} akun, {skipped} dilewati", ip=_client_ip(request),
    )
    return {"inserted": inserted, "skipped": skipped, "errors": errors[:50]}


@router.get("/accounts/template")
async def accounts_template(_: User = Depends(require_roles(*MASTER_WRITE_ROLES))):
    headers = ["code", "name", "category", "subcategory", "normal_balance"]
    rows = [["1.1.01.01", "Kas Bendahara", "aset", "kas_bank", "debit"]]
    blob = generate_excel_report(headers, rows, "BUMDES")
    return StreamingResponse(
        iter([blob]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Template-Kode-Akun.xlsx"'},
    )


@router.post("/transaction-types/import")
async def import_tx_types(
    request: Request,
    file: UploadFile = File(...),
    actor: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    content = await file.read()
    headers, data = parse_excel_rows(content)
    header = [h.strip().lower() for h in headers]
    required = {"code", "name", "debit", "credit"}
    if not required.issubset(set(header)):
        raise HTTPException(status_code=400, detail="Kolom wajib: code, name, debit, credit")
    inserted = skipped = 0
    for values in data:
        item = dict(zip(header, values))
        code = str(item.get("code") or "").strip()
        name = str(item.get("name") or "").strip()
        if not code or not name:
            continue
        group = str(item.get("group") or "BUMDES").strip().upper()
        exists = (
            await session.execute(
                select(TransactionType).where(TransactionType.code == code, TransactionType.group_code == group)
            )
        ).scalar_one_or_none()
        if exists:
            skipped += 1
            continue
        session.add(
            TransactionType(
                code=code,
                name=name,
                debit=str(item.get("debit") or "").strip(),
                credit=str(item.get("credit") or "").strip(),
                group_code=group,
            )
        )
        inserted += 1
    await record_audit(
        session, actor=actor, action="import_transaction_types", entity="transaction_types",
        detail=f"Import {inserted} jenis transaksi, {skipped} dilewati", ip=_client_ip(request),
    )
    return {"inserted": inserted, "skipped": skipped, "errors": []}


@router.get("/transaction-types/template")
async def tx_types_template(
    group: str = Query("BUMDES"),
    _: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
):
    headers = ["code", "name", "debit", "credit", "group"]
    rows = [["contoh_transaksi", "Contoh Transaksi", "1.1.01.01", "4.1.01.01", group]]
    blob = generate_excel_report(headers, rows, "Jenis Transaksi")
    return StreamingResponse(
        iter([blob]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Template-Jenis-Transaksi.xlsx"'},
    )


@router.get("/master-data/export")
async def export_master(
    group: str = Query("BUMDES"),
    section: str = Query("all"),
    _: User = Depends(require_roles(*MASTER_WRITE_ROLES)),
    session: AsyncSession = Depends(get_db),
):
    if section in {"all", "accounts"}:
        accs = (
            await session.execute(select(Account).where(Account.group_code == group).order_by(Account.code))
        ).scalars()
        headers = ["code", "name", "category", "subcategory", "normal_balance", "group"]
        rows = [[a.code, a.name, a.category, a.subcategory, a.normal_balance, a.group_code] for a in accs]
        blob = generate_excel_report(headers, rows, "Kode Akun")
        filename = f"Master-Data-Kode-Akun-{group}.xlsx"
    else:
        types = (
            await session.execute(
                select(TransactionType).where(TransactionType.group_code == group).order_by(TransactionType.code)
            )
        ).scalars()
        headers = ["code", "name", "debit", "credit", "group"]
        rows = [[t.code, t.name, t.debit, t.credit, t.group_code] for t in types]
        blob = generate_excel_report(headers, rows, "Jenis Transaksi")
        filename = f"Master-Data-Jenis-Transaksi-{group}.xlsx"
    return StreamingResponse(
        iter([blob]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
